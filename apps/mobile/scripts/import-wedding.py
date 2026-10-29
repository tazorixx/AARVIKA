"""Read-only source import. Run with the bundled Python/openpyxl runtime.
Writes only the generated mobile fixture; never publishes products or rituals.
"""
import json
import re
from pathlib import Path
import openpyxl

mobile = Path(__file__).resolve().parents[1]
workspace = mobile.parents[1]
source = workspace / 'references' / 'Gift4U_Indian_Wedding_Taxonomy_Research.xlsx'
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
def slug(text):
    return re.sub(r'[^a-z0-9]+', '-', text.lower()).strip('-')

traditions, ceremonies = [], []
for row in workbook['Community Templates'].iter_rows(min_row=2, values_only=True):
    if not row[0]:
        continue
    traditions.append(dict(id=slug(row[0]), name=row[0], faith=row[1], regions=[row[2]], ceremonyIds=[], reviewStatus='research-template', source=row[9] or 'Source workbook'))
for row_number, row in enumerate(workbook['Ceremony Catalogue'].iter_rows(min_row=2, values_only=True), 2):
    if not row[0]:
        continue
    tradition = next(t for t in traditions if t['name'] == row[0])
    ceremony_id = f"{tradition['id']}-{int(row[4])}-{slug(row[5])}"
    tradition['ceremonyIds'].append(ceremony_id)
    ceremonies.append(dict(id=ceremony_id, traditionId=tradition['id'], name=row[5], stage=row[3], optional=True, needs=[n.strip() for n in (row[6] or '').split(';') if n.strip()], reviewRule=row[7] or 'Family/officiant confirmation required', source=row[9] or '', sourceRange=f'Ceremony Catalogue!A{row_number}:J{row_number}'))
result = dict(source='Gift4U_Indian_Wedding_Taxonomy_Research.xlsx', status='research-only-not-reviewed', traditions=traditions, ceremonies=ceremonies)
(mobile / 'src/services/wedding-data.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f'Imported {len(traditions)} research traditions and {len(ceremonies)} optional ceremony slots.')
